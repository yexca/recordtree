from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
import pytest

from recordtree.app import RecordTreeApp
from recordtree.db import connect
from recordtree.exceptions import ValidationError


HEADERS = [
    "actor_raw",
    "delivery_date",
    "title",
    "entry_date",
    "note",
    "upload_title",
    "duplicate_search_raw",
    "source_name",
    "mega_json",
    "size_raw",
    "ignored_extra",
]

REAL_XLSX_HEADERS = [
    "声优",
    "配信日期",
    "标题",
    "录入日期",
    "备注",
    "上传标题",
    "重复检索",
    "来源",
    "MEGA",
    "容量",
    "ignored_extra",
]


def _mega(*links: dict[str, object]) -> str:
    return json.dumps(
        {
            "FileNames": "bundle",
            "total": sum(int(link["Size"]) for link in links),
            "FormattedSize": "3 KB",
            "property": list(links),
        },
        ensure_ascii=False,
    )


def _write_workbook(path: Path, headers: list[str] | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or HEADERS)
    sheet.append(
        [
            "Actor A",
            "2026-01-02",
            "Single link",
            "2026-01-03",
            "note",
            "Upload single",
            "",
            "Source A",
            _mega({"Link": "https://example.invalid/excel/1", "Size": 1024, "FormattedSize": "1 KB", "Type": "mp4"}),
            "1 KB",
            "extra",
        ]
    )
    sheet.append(
        [
            "Actor B",
            "2026-02-02",
            "Multi link",
            "2026-02-03",
            "",
            "Upload multi",
            "",
            "Source B",
            _mega(
                {"Link": "https://example.invalid/excel/2", "Size": 2048, "FormattedSize": "2 KB", "Type": ".mp4"},
                {"Link": "https://example.invalid/excel/3", "Size": 512, "FormattedSize": "512 B", "Type": ".srt"},
            ),
            "2.5 KB",
            "extra",
        ]
    )
    sheet.append(
        [
            "Actor C",
            "",
            "Blank dates",
            "",
            "",
            "Upload blank",
            "",
            "Source C",
            _mega({"Link": "https://example.invalid/excel/4", "Size": 4096, "FormattedSize": "4 KB", "Type": ""}),
            "",
            "extra",
        ]
    )
    sheet.append(
        [
            "Actor D",
            "2026-04-02",
            "Bad mega",
            "2026-04-03",
            "",
            "Upload bad",
            "",
            "Source D",
            "{not json",
            "1 KB",
            "extra",
        ]
    )
    workbook.save(path)


def _write_duplicate_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    base = [
        "Actor A",
        "2026-01-02",
        "Duplicate title",
        "2026-01-03",
        "note",
        "Duplicate upload",
        "",
        "Source A",
    ]
    sheet.append(
        base
        + [
            _mega({"Link": "https://example.invalid/duplicate/1", "Size": 1024, "FormattedSize": "1 KB", "Type": "mp4"}),
            "1 KB",
            "extra",
        ]
    )
    sheet.append(
        base
        + [
            _mega({"Link": "https://example.invalid/duplicate/2", "Size": 2048, "FormattedSize": "2 KB", "Type": "mp4"}),
            "2 KB",
            "extra",
        ]
    )
    workbook.save(path)


def test_excel_import_records_errors_and_is_idempotent(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    RecordTreeApp().init()
    workbook_path = tmp_path / "fixture.xlsx"
    _write_workbook(workbook_path)

    first = RecordTreeApp().import_file(workbook_path)
    second = RecordTreeApp().import_file(workbook_path)

    assert first.status == "completed_with_errors"
    assert first.stats.total_rows == 4
    assert first.stats.inserted_groups == 3
    assert first.stats.inserted_links == 4
    assert first.stats.error_count == 1
    assert first.error_csv_path is not None
    assert first.error_csv_path.exists()
    assert first.extra_columns == ("ignored_extra",)

    assert second.status == "completed_with_errors"
    assert second.stats.inserted_groups == 0
    assert second.stats.updated_groups == 0
    assert second.stats.unchanged_groups == 3
    assert second.stats.removed_groups == 0
    assert second.stats.error_count == 1

    conn = connect(tmp_path / "config" / "recordtree.sqlite3")
    try:
        assert conn.execute("SELECT COUNT(*) FROM record_groups").fetchone()[0] == 3
        assert conn.execute("SELECT COUNT(*) FROM download_links WHERE is_deleted = 0").fetchone()[0] == 4
        assert conn.execute("SELECT COUNT(*) FROM import_errors").fetchone()[0] == 2
        assert (
            conn.execute("SELECT status FROM imports WHERE id = ?", (first.import_id,)).fetchone()[0]
            == "completed_with_errors"
        )
    finally:
        conn.close()


def test_excel_import_accepts_real_chinese_headers(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    RecordTreeApp().init()
    workbook_path = tmp_path / "fixture.xlsx"
    _write_workbook(workbook_path, REAL_XLSX_HEADERS)

    result = RecordTreeApp().import_file(workbook_path)

    assert result.stats.total_rows == 4
    assert result.stats.inserted_groups == 3
    assert result.extra_columns == ("ignored_extra",)


def test_import_file_reports_progress(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    RecordTreeApp().init()
    workbook_path = tmp_path / "fixture.xlsx"
    _write_workbook(workbook_path, REAL_XLSX_HEADERS)
    events: list[tuple[str, int, int | None]] = []

    RecordTreeApp().import_file(
        workbook_path,
        progress_callback=lambda event: events.append((event.phase, event.completed_rows, event.total_rows)),
    )

    assert events == [
        ("Importing", 0, 4),
        ("Reading", 1, 4),
        ("Reading", 2, 4),
        ("Reading", 3, 4),
        ("Reading", 4, 4),
        ("Writing", 5, 7),
        ("Writing", 6, 7),
        ("Writing", 7, 7),
    ]


def test_excel_duplicate_source_keys_merge_links_and_reimport_is_idempotent(
    tmp_path: Path,
    monkeypatch,
) -> None:
    monkeypatch.chdir(tmp_path)
    RecordTreeApp().init()
    workbook_path = tmp_path / "duplicates.xlsx"
    _write_duplicate_workbook(workbook_path)

    first = RecordTreeApp().import_file(workbook_path)
    second = RecordTreeApp().import_file(workbook_path)

    assert first.stats.total_rows == 2
    assert first.stats.inserted_groups == 1
    assert first.stats.inserted_links == 2
    assert first.stats.link_sets_changed == 0

    assert second.stats.total_rows == 2
    assert second.stats.inserted_groups == 0
    assert second.stats.unchanged_groups == 1
    assert second.stats.link_sets_changed == 0
    assert second.stats.inserted_links == 0

    conn = connect(tmp_path / "config" / "recordtree.sqlite3")
    try:
        assert conn.execute("SELECT COUNT(*) FROM record_groups").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM download_links WHERE is_deleted = 0").fetchone()[0] == 2
        assert conn.execute("SELECT COUNT(*) FROM download_links WHERE is_deleted = 1").fetchone()[0] == 0
        notes = conn.execute("SELECT notes FROM imports WHERE id = ?", (first.import_id,)).fetchone()[0]
        assert "Duplicate Excel records merged: 1 source keys, 1 extra rows" in notes
    finally:
        conn.close()


def test_excel_import_reads_file_names_and_raw_line_breaks(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    RecordTreeApp().init()
    workbook_path = tmp_path / "current.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(REAL_XLSX_HEADERS[:-1])
    name = "[Actor A][2026-01-02][Title][2026-01-03]"
    mega = (
        '{"FileNames": "' + name + '", "total": 0, "FormattedSize": "0 KB", "property": ['
        '{"Link": "https://example.invalid/current/1", "Size": 2048, "FormattedSize": "2 KB", "Type": "' + name + '.mp4"},'
        '{"Link": "https://example.invalid/current/2", "Size": 64, "FormattedSize": "64 B", "Type": ".exe\r\n"}]}'
    )
    sheet.append(["Actor A", "2026-01-02", "Title", "2026-01-03", None, name, "46024 - Actor A", "Withny", mega, "2 KB"])
    workbook.save(workbook_path)

    result = RecordTreeApp().import_file(workbook_path)

    assert result.status == "completed"
    conn = connect(tmp_path / "config" / "recordtree.sqlite3")
    try:
        rows = conn.execute("SELECT file_name, file_type FROM download_links ORDER BY link_order").fetchall()
    finally:
        conn.close()
    assert [tuple(row) for row in rows] == [(f"{name}.mp4", ".mp4"), (None, ".exe")]


def test_workbook_without_valid_rows_leaves_the_library_unchanged(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    RecordTreeApp().init()
    full = tmp_path / "full.xlsx"
    _write_workbook(full)
    RecordTreeApp().import_file(full)
    empty = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.active.append(HEADERS)
    workbook.save(empty)

    with pytest.raises(ValidationError, match="no valid rows"):
        RecordTreeApp().import_file(empty)

    conn = connect(tmp_path / "config" / "recordtree.sqlite3")
    try:
        assert conn.execute("SELECT COUNT(*) FROM record_groups WHERE is_deleted = 0").fetchone()[0] == 3
        assert conn.execute("SELECT status FROM imports ORDER BY id DESC").fetchone()[0] == "failed"
    finally:
        conn.close()


def test_excel_row_with_null_mega_imports_without_links(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    RecordTreeApp().init()
    workbook_path = tmp_path / "pending.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(REAL_XLSX_HEADERS[:-1])
    sheet.append(["Actor A", "2026-01-02", "Pending upload", "2026-01-03", None, "Upload", None, "Withny", "null", "65 MB"])
    workbook.save(workbook_path)

    result = RecordTreeApp().import_file(workbook_path)

    assert result.status == "completed"
    assert result.stats.error_count == 0
    assert result.stats.inserted_groups == 1
    assert result.stats.inserted_links == 0
    detail = RecordTreeApp().info("1")
    assert detail.title == "Pending upload"
    assert detail.links == []
    assert detail.downloaded == "unknown"
